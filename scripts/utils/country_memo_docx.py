"""Word (.docx) files of the Country Insights policy memos (spec_country.md B7.4).

Run with `python -I` by scripts/88_export_country_data.R --render-memos (scripts/utils/country_memos.R):

    python -I country_memo_docx.py --in memo.ast.json --out memo.docx      build the file from the memo AST
    python -I country_memo_docx.py --check memo.docx --ast memo.ast.json   read it back; exit 1 unless the
                                                                           paragraph texts equal the AST's
    python -I country_memo_docx.py --pdf-text memo.pdf [--out text.json]   print (or write) {"pages": n, "text": ...}

The AST is the memo document of R/dashboard/country/country_memo_lint.R (country_memo_document()): the
same structure the screen and the PDF are rendered from, so the three carry the same words. US Letter,
0.75 in margins, Arial 10.5 pt (Word may not have Roboto); the summary as a one-cell table shaded F3F4F6
with an orange (F58229) left border; numbered Heading 1 sections; List Bullet items; superscript source
numbers; the source list as List Number paragraphs; the label in the footer. Needs python-docx >= 1.1
and pypdf.
"""

import argparse
import json
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x4B, 0x55, 0x63)


def item_text(item):
    """The paragraph text Word reads back: the inlines, then the source numbers joined by commas."""
    text = "".join(n["v"] for n in item["inlines"])
    refs = item.get("refs") or []
    return text + (",".join(str(r) for r in refs) if refs else "")


def expected_texts(doc):
    """Body paragraph texts in reading order (the same rule as country_memo_texts() in R)."""
    out = [doc["title"], doc["subtitle"], doc["label"], "Summary"]
    out += [item_text(it) for it in doc["summary"]]
    for s in doc["sections"]:
        out.append(f"{s['n']}. {s['title']}")
        for b in s["blocks"]:
            if b["type"] == "ul":
                out += [item_text(it) for it in b["items"]]
            else:
                out.append(item_text(b))
    out.append("Sources")
    out += [s["text"] for s in doc["sources"]]
    out.append(doc["label"])
    return out


def add_runs(par, item, size=None):
    for n in item["inlines"]:
        run = par.add_run(n["v"])
        if n["t"] == "strong":
            run.bold = True
        elif n["t"] == "em":
            run.italic = True
        if size:
            run.font.size = size
    refs = item.get("refs") or []
    if refs:
        sup = par.add_run(",".join(str(r) for r in refs))
        sup.font.superscript = True
        sup.font.color.rgb = MUTED


def shade_cell(cell, fill="F3F4F6", border="F58229"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        if side == "left":
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "24")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), border)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tc_pr.append(borders)
    mar = OxmlElement("w:tcMar")
    for side, w in (("top", 120), ("left", 200), ("bottom", 80), ("right", 160)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def hairline_above(par, color="D1D5DB"):
    p_pr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")
    top.set(qn("w:space"), "4")
    top.set(qn("w:color"), color)
    bdr.append(top)
    p_pr.append(bdr)


def build(doc, out):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.75))
    normal = d.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for name in ("Title", "Heading 1", "Heading 2", "List Bullet", "List Number"):
        st = d.styles[name]
        st.font.name = "Arial"
        st.font.color.rgb = INK
    d.styles["Title"].font.size = Pt(20)
    d.styles["Heading 1"].font.size = Pt(13)
    d.styles["Heading 2"].font.size = Pt(10)

    cp = d.core_properties
    cp.title = doc["title"]
    cp.subject = "Country Insights policy memo"
    cp.author = "NEIS Center, Allied Industrial Base Explorer (AI-generated, checked against the data)"
    cp.keywords = "energy security; economic opportunity; Strategic Index; Partnership Strength Index; " + doc["iso3"]
    cp.comments = f"paper baseline as of {doc.get('as_of_label') or doc.get('as_of')}; fact sheet sha256 {doc.get('facts_sha256', '')}"

    d.add_paragraph(doc["title"], style="Title")
    sub = d.add_paragraph()
    r = sub.add_run(doc["subtitle"])
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    lab = d.add_paragraph()
    r = lab.add_run(doc["label"])
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = MUTED
    hairline_above(lab)

    table = d.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    shade_cell(cell)
    head = cell.paragraphs[0]
    hr = head.add_run("Summary")
    hr.bold = True
    hr.font.size = Pt(12)
    for it in doc["summary"]:
        add_runs(cell.add_paragraph(style="List Bullet"), it)
    d.add_paragraph()  # space after the box (empty: skipped by the read-back)

    for s in doc["sections"]:
        d.add_paragraph(f"{s['n']}. {s['title']}", style="Heading 1")
        for b in s["blocks"]:
            if b["type"] == "ul":
                for it in b["items"]:
                    add_runs(d.add_paragraph(style="List Bullet"), it)
            else:
                add_runs(d.add_paragraph(), b)

    src_head = d.add_paragraph("Sources", style="Heading 2")
    hairline_above(src_head)
    for s in doc["sources"]:
        p = d.add_paragraph(style="List Number")
        r = p.add_run(s["text"])
        r.font.size = Pt(8.5)
    end = d.add_paragraph()
    r = end.add_run(doc["label"])
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    hairline_above(end)

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fr = foot.add_run(doc["label"] + (f"  {doc['footer']}" if doc.get("footer") else ""))
    fr.italic = True
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUTED
    d.save(out)


def body_texts(path):
    """Paragraph texts of the body in document order, table cells included; empty paragraphs skipped."""
    d = Document(path)
    out = []
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            t = "".join(x.text or "" for x in el.iter(qn("w:t")))
            if t.strip():
                out.append(t)
        elif tag == "tbl":
            for p in el.iter(qn("w:p")):
                t = "".join(x.text or "" for x in p.iter(qn("w:t")))
                if t.strip():
                    out.append(t)
    return out


def check(path, ast_path):
    with open(ast_path, encoding="utf-8") as f:
        doc = json.load(f)
    want = expected_texts(doc)
    got = body_texts(path)
    if got == want:
        print(f"ok: {len(got)} paragraphs")
        return 0
    n = min(len(got), len(want))
    for i in range(n):
        if got[i] != want[i]:
            print(f"paragraph {i + 1} differs:\n  docx: {got[i][:200]}\n  ast:  {want[i][:200]}")
            break
    else:
        print(f"{len(got)} paragraphs in the file, {len(want)} expected")
    return 1


def pdf_text(path, out=None):
    from pypdf import PdfReader

    reader = PdfReader(path)
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    res = {"pages": len(reader.pages), "text": text}
    # to a file when asked: R reads a long single-line stdout in pieces of about 8 kB (system2)
    if out:
        with open(out, "w", encoding="utf-8") as f:
            json.dump(res, f, ensure_ascii=True)
    else:
        json.dump(res, sys.stdout, ensure_ascii=True)
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--in", dest="inp")
    ap.add_argument("--out")
    ap.add_argument("--check")
    ap.add_argument("--ast")
    ap.add_argument("--pdf-text")
    a = ap.parse_args()
    if a.pdf_text:
        return pdf_text(a.pdf_text, a.out)
    if a.check:
        return check(a.check, a.ast)
    if a.inp and a.out:
        with open(a.inp, encoding="utf-8") as f:
            doc = json.load(f)
        build(doc, a.out)
        return 0
    ap.print_help()
    return 2


if __name__ == "__main__":
    sys.exit(main())
